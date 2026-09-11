#!/usr/bin/env python3
"""Convert a client's terminology spreadsheet into a TermGuard rulebook.

    python scripts/import_rulebook.py terms.xlsx [-o data/rulebook.yaml]

Every terminology programme already has this spreadsheet. It is usually three columns -
old term, new term, a notes field - and the notes are where the judgment lives: "only in
clinical sections", "depends on audience", "except in quoted regulation". Rows whose notes
say something like that become ``context_required`` rules, so they reach a human instead
of being applied blindly, with the note carried across as the reviewer's guidance.

Accepts .xlsx (needs openpyxl) or .csv. Rows that cannot be interpreted are printed rather
than dropped silently, because a rule missing from the book is a violation nobody will
ever see.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path
from typing import Any, Iterator, Sequence

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

from termguard.rulebook import CaseKind, MatchKind, Rule, Rulebook, dump_rulebook  # noqa: E402

# Header spellings seen in the wild, normalized.
OLD_HEADERS = {"old term", "old", "deprecated", "deprecated term", "from", "avoid", "do not use"}
NEW_HEADERS = {"new term", "new", "approved", "approved term", "to", "use", "preferred"}
NOTE_HEADERS = {"notes", "note", "comment", "comments", "guidance", "rationale"}

# A note containing any of these is a judgment call, not a substitution.
CONTEXT_MARKERS = ("context", "depends", "except", "only", "unless", "sometimes",
                   "case by case", "case-by-case", "varies", "audience")


def normalize(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def read_rows(path: Path) -> tuple[list[str], list[list[str]]]:
    """Return (header, rows) from .xlsx or .csv."""
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError:  # pragma: no cover - depends on the optional extra
            raise SystemExit(
                "reading .xlsx needs openpyxl:\n"
                "    pip install openpyxl\n"
                "or export the sheet as .csv and pass that instead"
            )
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = [[normalize(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
        workbook.close()
    elif path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as handle:
            rows = [[normalize(cell) for cell in row] for row in csv.reader(handle)]
    else:
        raise SystemExit(f"unsupported file type: {path.suffix} (want .xlsx or .csv)")

    rows = [row for row in rows if any(row)]
    if not rows:
        raise SystemExit(f"{path} has no data")
    return rows[0], rows[1:]


def locate_columns(header: Sequence[str]) -> tuple[int, int, int | None]:
    """Find the old/new/notes columns, however they happen to be spelled."""
    lowered = [h.casefold() for h in header]

    def find(candidates: set[str]) -> int | None:
        for index, name in enumerate(lowered):
            if name in candidates:
                return index
        for index, name in enumerate(lowered):
            if any(candidate in name for candidate in candidates):
                return index
        return None

    old, new, note = find(OLD_HEADERS), find(NEW_HEADERS), find(NOTE_HEADERS)
    if old is None or new is None:
        raise SystemExit(
            "could not find the old-term and new-term columns.\n"
            f"  header row: {list(header)}\n"
            f"  expected something like: {sorted(OLD_HEADERS)[:3]} / {sorted(NEW_HEADERS)[:3]}"
        )
    return old, new, note


def needs_context(note: str) -> bool:
    lowered = note.casefold()
    return any(marker in lowered for marker in CONTEXT_MARKERS)


def build_rule(index: int, old: str, new: str, note: str, owner: str,
               *, case_only: bool = False) -> Rule:
    """One spreadsheet row as a rule.

    Defaults are deliberately conservative: whole-word matching so a term never fires
    inside a longer word, and case preservation so heading capitalization survives.
    Multi-word terms become phrase matches, which tolerate the extra whitespace Word
    leaves behind.
    """
    contextual = needs_context(note)
    return Rule(
        id=f"R-{index:03d}",
        deprecated=[old],
        approved=new,
        match=MatchKind.PHRASE if " " in old else MatchKind.WHOLE_WORD,
        # A case-only rule must replace verbatim; preserving the source's casing would
        # reproduce exactly the spelling the rule exists to fix.
        case=CaseKind.EXACT if case_only else CaseKind.PRESERVE,
        context_required=contextual,
        context_note=note if contextual else "",
        rationale=note if not contextual else "",
        owner=owner,
    )


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("spreadsheet", type=Path)
    parser.add_argument("-o", "--output", type=Path, default=REPO_ROOT / "data" / "rulebook.yaml")
    parser.add_argument("--owner", default="imported", help="owner recorded on every rule")
    parser.add_argument("--dry-run", action="store_true", help="report only; write nothing")
    args = parser.parse_args()

    header, rows = read_rows(args.spreadsheet)
    old_col, new_col, note_col = locate_columns(header)

    rules: list[Rule] = []
    skipped: list[tuple[int, str, str]] = []
    seen: dict[str, int] = {}

    for offset, row in enumerate(rows, start=2):  # start=2: row 1 is the header
        def cell(index: int | None) -> str:
            return row[index] if index is not None and index < len(row) else ""

        old, new, note = cell(old_col), cell(new_col), cell(note_col)

        if not old or not new:
            skipped.append((offset, f"{old} -> {new}", "old or new term is blank"))
            continue
        if old == new:
            skipped.append((offset, old, "old and new term are identical"))
            continue
        if old.casefold() == new.casefold():
            # A case-only change is a real rule ("ml" -> "mL"), but preserving the
            # source's casing would undo it, so such rules must replace verbatim.
            case_only = True
        else:
            case_only = False
        if old.casefold() in seen:
            skipped.append((offset, old, f"duplicate of row {seen[old.casefold()]}"))
            continue

        try:
            rules.append(build_rule(len(rules) + 1, old, new, note, args.owner,
                                    case_only=case_only))
        except Exception as exc:  # noqa: BLE001 - report, do not abort the whole import
            skipped.append((offset, old, f"invalid: {exc}"))
            continue
        seen[old.casefold()] = offset

    contextual = [r for r in rules if r.context_required]
    print(f"\nread {len(rows)} rows from {args.spreadsheet.name}")
    print(f"  columns: old={header[old_col]!r} new={header[new_col]!r} "
          f"notes={header[note_col]!r}" if note_col is not None else
          f"  columns: old={header[old_col]!r} new={header[new_col]!r} notes=(none)")
    print(f"  {len(rules)} rules built")
    print(f"  {len(contextual)} marked context_required (their notes mention "
          f"context, exceptions or conditions)")
    for rule in contextual[:8]:
        print(f"      {rule.id}  {rule.deprecated[0]!r} -> {rule.approved!r}")
        print(f"            note: {rule.context_note[:88]}")

    if skipped:
        print(f"\n  {len(skipped)} row(s) could not be interpreted:")
        for line, term, why in skipped:
            print(f"      row {line}: {term!r} - {why}")

    if not rules:
        print("\nnothing to write.")
        return 1

    if args.dry_run:
        print("\ndry run: nothing written.")
        return 0

    book = Rulebook(rules=rules, version="1")
    dump_rulebook(book, args.output)
    print(f"\nwrote {args.output} ({len(rules)} rules)")
    print("\nReview before using it. The importer guesses conservatively; it cannot know")
    print("which terms need an exception for quoted regulatory text, or which rules should")
    print("be scoped to particular document parts.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
