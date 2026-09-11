#!/usr/bin/env python3
"""Generate the synthetic Meridian Medical corpus with planted, recorded violations.

Real client files cannot leave the building, so the demo needs documents whose every
terminology violation is known in advance. This script writes 24 documents with planted
violations plus 2 clean controls, and records the ground truth so the scanner's recall and
precision can be *measured* rather than asserted.

Violations are deliberately placed where manual review misses them: headers, footers,
table cells, footnotes, and Title Case headings. Three kinds are planted:

    change   a genuine violation the deterministic pipeline must fix
    keep     a match that an exception protects (quoted CFR text, historical name)
    judge    a context_required match that must reach a human, not be auto-changed

Run: python scripts/make_corpus.py
"""

from __future__ import annotations

import json
import random
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from footnotes import inject_footnotes  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parent.parent
CORPUS = REPO_ROOT / "data" / "corpus"
SEED = 20240301


@dataclass(frozen=True)
class V:
    """A planted violation inside a text segment."""

    text: str
    rule: str
    action: str  # change | keep | judge


Segment = str | V


class DocBuilder:
    """Wraps python-docx and records ground truth as content is written."""

    def __init__(self, name: str, doc_number: str, doc_type: str, revision: str) -> None:
        self.name = name
        self.doc_number = doc_number
        self.doc_type = doc_type
        self.doc = docx.Document()
        self.truth: list[dict] = []
        self._body_index = 0
        self._footnote_markers: dict[str, str] = {}
        self._fn_counter = 0

        for style_name, size in (("Normal", 11),):
            self.doc.styles[style_name].font.size = Pt(size)

        section = self.doc.sections[0]
        section.header.paragraphs[0].text = ""
        section.footer.paragraphs[0].text = ""
        self._section = section

    # -- recording -----------------------------------------------------------

    def _render(
        self,
        segments: Sequence[Segment],
        *,
        part: str,
        paragraph_index: int,
        in_table: bool = False,
        is_heading: bool = False,
        container: str | None = None,
    ) -> str:
        """Join segments into text and record every violation with its offset."""
        text = ""
        counts: dict[str, int] = {}
        for seg in segments:
            if isinstance(seg, V):
                start = len(text)
                key = seg.text.casefold()
                occurrence = counts.get(key, 0)
                counts[key] = occurrence + 1
                self.truth.append(
                    {
                        "file": self.name,
                        "part": part,
                        "container": container,
                        "in_table": in_table,
                        "is_heading": is_heading,
                        "paragraph_index": paragraph_index,
                        "text": seg.text,
                        "rule_id": seg.rule,
                        "expected_action": seg.action,
                        "span_start": start,
                        "occurrence": occurrence,
                    }
                )
                text += seg.text
            else:
                text += seg
        return text

    # -- content -------------------------------------------------------------

    def header(self, *segments: Segment) -> None:
        para = self._section.header.paragraphs[0]
        para.text = self._render(segments, part="header", paragraph_index=0)

    def footer(self, *segments: Segment) -> None:
        para = self._section.footer.paragraphs[0]
        para.text = self._render(segments, part="footer", paragraph_index=0)

    def heading(self, *segments: Segment, level: int = 1) -> None:
        text = self._render(
            segments, part="body", paragraph_index=self._body_index, is_heading=True
        )
        self.doc.add_heading(text, level=level)
        self._body_index += 1

    def para(self, *segments: Segment) -> None:
        text = self._render(segments, part="body", paragraph_index=self._body_index)
        self.doc.add_paragraph(text)
        self._body_index += 1

    def bullet(self, *segments: Segment) -> None:
        text = self._render(segments, part="body", paragraph_index=self._body_index)
        self.doc.add_paragraph(text, style="List Bullet")
        self._body_index += 1

    def para_with_footnote(self, body: Sequence[Segment], note: Sequence[Segment]) -> None:
        """A body paragraph carrying a footnote. Violations may sit in either."""
        self._fn_counter += 1
        marker = f"[[FN{self._fn_counter}]]"
        text = self._render(body, part="body", paragraph_index=self._body_index)
        note_text = self._render(
            note, part="footnote", paragraph_index=self._fn_counter - 1
        )
        self.doc.add_paragraph(text + marker)
        self._footnote_markers[marker] = note_text
        self._body_index += 1

    def table(self, rows: Sequence[Sequence[Sequence[Segment]]], *, index: int = 1) -> None:
        """A table of cells; each cell is a sequence of segments."""
        n_rows, n_cols = len(rows), max(len(r) for r in rows)
        table = self.doc.add_table(rows=n_rows, cols=n_cols)
        table.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, cell_segments in enumerate(row):
                container = f"table {index} / row {r + 1} / cell {c + 1}"
                text = self._render(
                    cell_segments,
                    part="body",
                    paragraph_index=self._body_index,
                    in_table=True,
                    container=container,
                )
                table.cell(r, c).text = text
                self._body_index += 1

    def page_break(self) -> None:
        self.doc.add_page_break()

    # -- output --------------------------------------------------------------

    def save(self, directory: Path) -> Path:
        path = directory / self.name
        self.doc.save(path)
        if self._footnote_markers:
            inject_footnotes(path, self._footnote_markers)
        return path


# --------------------------------------------------------------------- content

def build_ifu(n: int) -> DocBuilder:
    """Instructions for Use excerpt. Carries the patient-facing plain-language sections."""
    b = DocBuilder(f"IFU-{n:03d}.docx", f"IFU-{n:03d}", "ifu", "Rev C")
    b.header("Document ", V("IFU", "R-003", "keep"), f"-{n:03d}", " - ",
             V("Meridian Pump 2", "R-001", "change"), " Instructions for Use")
    b.footer("Meridian Medical - Rev C - ", V("single use", "R-008", "change"), " device")

    b.heading("Instructions for Use: ", V("Meridian Pump 2", "R-001", "change"), level=1)
    b.para("This document describes safe operation of the device for the trained user. ",
           "Read all warnings before first use.")

    b.heading("Intended User", level=2)
    b.para("The device is intended for use by a ", V("physician", "R-004", "change"),
           " or other trained clinician in a hospital setting. Refer to the ",
           V("IFU", "R-003", "judge"),
           " for cleaning instructions before first use.")
    b.para("Operators ", V("shall", "R-005", "change"),
           " complete the training module before independent use.")

    b.heading("Setting Up The ", V("Infusion Set", "R-009", "change"), level=2)
    b.para("Attach the ", V("infusion set", "R-009", "change"),
           " to the pump and prime the line. Each ", V("infusion set", "R-009", "change"),
           " is ", V("single use", "R-008", "change"), " only.")
    b.bullet("Confirm the reservoir volume does not exceed 250 ",
             V("ml", "R-007", "change"), ".")
    b.bullet("Select the protocol from the medication library before starting therapy.")

    b.heading("What You May Experience", level=2)
    b.para("Talk to your care team if you notice any ", V("side effect", "R-002", "judge"),
           " while using the pump. Common ", V("side effects", "R-002", "judge"),
           " include redness at the infusion site. Most ",
           V("side effects", "R-002", "judge"), " are mild and go away on their own.")

    b.heading("Regulatory Note", level=2)
    b.para('The regulation states that ', V('"the physician shall maintain records"', "R-004", "keep"),
           " per 21 CFR 820.180.")

    b.table([
        [["Parameter"], ["Limit"]],
        [["Maximum rate"], ["999 ", V("ml", "R-007", "change"), "/hr"]],
        [[V("Infusion Set", "R-009", "change"), " life"], ["96 hours, ",
                                                           V("single use", "R-008", "change")]],
        [["Alarm condition display"], ["Front panel"]],
    ])
    return b


def build_sop(n: int) -> DocBuilder:
    """Standard operating procedure page. Heavy on 'shall' and process language."""
    b = DocBuilder(f"SOP-{n:03d}.docx", f"SOP-{n:03d}", "sop", "Rev 4")
    b.header("SOP-", f"{n:03d}", " - Device Servicing - ", V("Meridian Pump 2", "R-001", "change"))
    b.footer("Controlled copy - Rev 4 - page 1 of 2")

    b.heading("Standard Operating Procedure: Device Servicing", level=1)
    b.para("This procedure applies to all servicing activity on the ",
           V("Meridian Pump 2", "R-001", "change"), " platform.")

    b.heading("Responsibilities", level=2)
    b.para("The service technician ", V("shall", "R-005", "change"),
           " record each intervention in the maintenance log. The supervising ",
           V("physician", "R-004", "change"), " reviews the log monthly.")
    b.para("Any ", V("error message", "R-011", "change"),
           " observed during servicing ", V("shall", "R-005", "change"),
           " be recorded with its code and timestamp.")

    b.heading("Procedure", level=2)
    b.bullet("Disconnect the ", V("infusion set", "R-009", "change"), " before opening the case.")
    b.bullet("Verify the ", V("drug library", "R-012", "change"), " version matches the release note.")
    b.bullet("Confirm the reservoir is drained below 5 ", V("ml", "R-007", "change"), ".")

    b.heading("Historical Note", level=2)
    b.para("This procedure supersedes SVC-11, which covered the device ",
           V("formerly known as the Meridian Pump 2", "R-001", "keep"),
           ", and is retained for traceability.")

    b.table([
        [["Step"], ["Owner"], ["Record"]],
        [["Inspect ", V("infusion set", "R-009", "change"), " port"], ["Technician"], ["Form S-1"]],
        [["Review ", V("error message", "R-011", "change"), " log"], [V("Nurse", "R-010", "judge")],
         ["Form S-2"]],
    ])
    return b


def build_risk(n: int) -> DocBuilder:
    """Risk management summary. Clinical register wording; 'side effect' is wrong here."""
    b = DocBuilder(f"RMS-{n:03d}.docx", f"RMS-{n:03d}", "risk", "Rev B")
    b.header("Risk Management Summary ", f"RMS-{n:03d}", " - ", V("labelling", "R-006", "change"),
             " controls")
    b.footer("Meridian Medical - Rev B - confidential")

    b.heading("Risk Management Summary", level=1)
    b.para("This summary records residual risk for the ",
           V("Meridian Pump 2", "R-001", "change"), " following design verification.")

    b.heading("Residual Risk", level=2)
    b.para("Each identified ", V("side effect", "R-002", "judge"),
           " has been evaluated against the acceptability criteria in the risk plan. "
           "Occlusion remains the highest-severity hazard.")
    b.para("Mitigations rely on ", V("labelling", "R-006", "change"),
           " and on the alarm subsystem. The user ", V("shall", "R-005", "change"),
           " respond to every alarm condition within 30 seconds.")

    b.para_with_footnote(
        body=["Residual risk is accepted for the ", V("infusion set", "R-009", "change"),
              " occlusion hazard."],
        note=["Occlusion testing was performed at 250 ", V("ml", "R-007", "change"),
              "/hr by the supervising ", V("physician", "R-004", "change"), "."],
    )

    b.table([
        [["Hazard"], ["Severity"], ["Mitigation"]],
        [["Over-infusion"], ["Critical"], ["Medication library limits"]],
        [["Occlusion"], ["Serious"], [V("Error message", "R-011", "change"), " and alarm"]],
    ])
    return b


def build_clinical(n: int) -> DocBuilder:
    """Clinical evaluation excerpt. 'side effect' here is a genuine violation to judge."""
    b = DocBuilder(f"CER-{n:03d}.docx", f"CER-{n:03d}", "clinical", "Rev A")
    b.header("Clinical Evaluation ", f"CER-{n:03d}", " - ", V("Meridian Pump 2", "R-001", "change"))
    b.footer("Rev A - ", V("labelling", "R-006", "change"), " support document")

    b.heading("Clinical Evaluation Report Excerpt", level=1)
    b.para("Data were collected across four sites between 2023 and 2024.")

    b.heading("Reported Events", level=2)
    b.para("Investigators recorded each ", V("side effect", "R-002", "judge"),
           " observed during the study period. The reporting ",
           V("physician", "R-004", "change"),
           " classified events by severity and relatedness.")
    b.para("Device-related events were reviewed by the ", V("nurse", "R-010", "judge"),
           " coordinator and the medical monitor.")

    b.para_with_footnote(
        body=["Enrolment was limited to sites using the standard ",
              V("infusion set", "R-009", "change"), "."],
        note=["Site 3 reported one ", V("side effect", "R-002", "judge"),
              " requiring device replacement."],
    )

    b.table([
        [["Site"], ["Events"], ["Reviewer"]],
        [["Site 1"], ["4"], [V("Physician", "R-004", "change"), " monitor"]],
        [["Site 2"], ["2"], ["Study coordinator"]],
    ])
    return b


def build_labeling(n: int) -> DocBuilder:
    """Labeling draft. Mixes patient-facing plain language with regulated wording."""
    b = DocBuilder(f"LBL-{n:03d}.docx", f"LBL-{n:03d}", "labeling", "Rev 2")
    b.header("Labeling draft ", f"LBL-{n:03d}", " - ", V("single use", "R-008", "change"),
             " accessory")
    b.footer("Rev 2 - ", V("Meridian Pump 2", "R-001", "change"), " family")

    b.heading("Carton Text And ", V("Labelling", "R-006", "change"), level=1)
    b.para("Text below is set in 8pt and reviewed for readability.")

    b.heading("For The Patient", level=2)
    b.para("Tell your care team about any ", V("side effect", "R-002", "judge"),
           " you notice. Do not reuse the ", V("infusion set", "R-009", "change"),
           "; it is ", V("single use", "R-008", "change"), " only.")

    b.heading("For The Clinician", level=2)
    b.para("The ", V("physician", "R-004", "change"),
           " or prescribing provider configures the medication library before dispatch.")
    b.para("Reservoir capacity is 500 ", V("ml", "R-007", "change"), ".")

    b.para('Quoted directly from the predicate submission: ',
           V('"the nurse shall confirm the dose"', "R-010", "keep"), ".")

    b.table([
        [["Panel"], ["Text"]],
        [["Front"], [V("Meridian Pump 2", "R-001", "change")]],
        [["Back"], [V("Single use", "R-008", "change"), " only - 500 ",
                    V("ml", "R-007", "change")]],
    ])
    return b


def build_control(n: int) -> DocBuilder:
    """A clean document. Nothing here may ever be flagged."""
    b = DocBuilder(f"CTL-{n:03d}.docx", f"CTL-{n:03d}", "control", "Rev 1")
    b.header("Control document CTL-", f"{n:03d}", " - Meridian Infusion System")
    b.footer("Meridian Medical - Rev 1 - clean control")

    b.heading("Control Document: Meridian Infusion System", level=1)
    b.para("This document contains only approved terminology and must never be flagged "
           "by the scanner. It exercises the false-positive path.")
    b.heading("Approved Usage", level=2)
    b.para("The healthcare provider configures the medication library before dispatch. "
           "Each administration set is single-use only and must be discarded after 96 hours.")
    b.para("Reservoir capacity is 500 mL. Any alarm condition must be recorded in the "
           "maintenance log, and the clinician must respond within 30 seconds.")
    b.para("Adverse events are reported per the labeling and the risk management plan.")
    b.table([
        [["Parameter"], ["Limit"]],
        [["Maximum rate"], ["999 mL/hr"]],
        [["Administration set life"], ["96 hours, single-use"]],
    ])
    return b


BUILDERS: list[tuple[str, int, callable]] = [
    ("ifu", 6, build_ifu),
    ("sop", 6, build_sop),
    ("risk", 4, build_risk),
    ("clinical", 4, build_clinical),
    ("labeling", 4, build_labeling),
]


def main() -> int:
    random.seed(SEED)
    if CORPUS.exists():
        shutil.rmtree(CORPUS)
    CORPUS.mkdir(parents=True, exist_ok=True)

    truth: list[dict] = []
    written: list[str] = []

    for _kind, count, builder in BUILDERS:
        for i in range(1, count + 1):
            b = builder(i)
            b.save(CORPUS)
            truth.extend(b.truth)
            written.append(b.name)

    for i in range(1, 3):
        b = build_control(i)
        b.save(CORPUS)
        truth.extend(b.truth)  # empty by construction
        written.append(b.name)

    (CORPUS / "ground_truth.json").write_text(
        json.dumps(
            {
                "seed": SEED,
                "documents": written,
                "items": truth,
            },
            indent=2,
        )
    )

    # -- report against the corpus requirements ------------------------------
    def count(pred) -> int:
        return sum(1 for t in truth if pred(t))

    checks = {
        "documents": (len(written), 26),
        "violations total": (len(truth), None),
        "in headers/footers": (count(lambda t: t["part"] in {"header", "footer"}), 5),
        "in table cells": (count(lambda t: t["in_table"]), 6),
        "in footnotes": (count(lambda t: t["part"] == "footnote"), 3),
        "in headings": (count(lambda t: t["is_heading"]), 4),
        "exception cases (keep)": (count(lambda t: t["expected_action"] == "keep"), 3),
        "context cases (judge)": (count(lambda t: t["expected_action"] == "judge"), 8),
        "deterministic (change)": (count(lambda t: t["expected_action"] == "change"), None),
    }
    print(f"wrote {len(written)} documents to {CORPUS.relative_to(REPO_ROOT)}\n")
    ok = True
    for label, (actual, minimum) in checks.items():
        if minimum is None:
            print(f"  {label:26} {actual}")
        else:
            good = actual >= minimum
            ok &= good
            print(f"  {label:26} {actual:4}  (min {minimum}) {'OK' if good else 'SHORT'}")

    per_doc = {}
    for t in truth:
        per_doc.setdefault(t["file"], set()).add(t["rule_id"])
    violating = {f: len(r) for f, r in per_doc.items()}
    bad = {f: n for f, n in violating.items() if not (3 <= n <= 8)}
    docs_with_violations = [w for w in written if not w.startswith("CTL")]
    missing = [d for d in docs_with_violations if d not in per_doc]
    print(f"\n  rules violated per document: min={min(violating.values())} "
          f"max={max(violating.values())}")
    if bad:
        print(f"  OUT OF RANGE (want 3-8): {bad}")
        ok = False
    if missing:
        print(f"  DOCUMENTS WITH NO VIOLATIONS: {missing}")
        ok = False
    controls = [t for t in truth if t["file"].startswith("CTL")]
    if controls:
        print(f"  CONTROLS ARE NOT CLEAN: {len(controls)} planted items")
        ok = False
    else:
        print("  controls clean: yes (2 files, 0 planted items)")

    print("\n" + ("all corpus requirements met" if ok else "REQUIREMENTS NOT MET"))
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
